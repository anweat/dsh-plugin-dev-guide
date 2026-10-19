"""生成软件说明书 Word 版（软著申请用）。"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = r"C:\Users\zhito\WorkBuddy\2026-08-17-09-25-53\软件说明书.docx"

TITLE_FONT = "黑体"
BODY_FONT = "宋体"


def set_run(run, size=12, bold=False, cn=BODY_FONT, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if color:
        run.font.color.rgb = RGBColor(*color)


def h(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    size = {1: 16, 2: 14, 3: 13}[level]
    set_run(run, size=size, bold=True, cn=TITLE_FONT)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.style = doc.styles[f"Heading {level}"]
    # 覆盖标题样式字体
    for r in p.runs:
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), TITLE_FONT)
        r.font.color.rgb = RGBColor(0, 0, 0)
        r.font.bold = True
        r.font.size = Pt(size)
    return p


def para(doc, text, indent=True, size=12, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=size, bold=bold)
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.line_spacing = 1.5
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run(run, size=12)
    p.paragraph_format.line_spacing = 1.5
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        set_run(run, size=10.5, bold=True, cn=TITLE_FONT)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9E2F3")
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            set_run(run, size=10.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_txt = OxmlElement("w:t")
    fld_txt.text = "（打开文档后，在目录处右键选择「更新域」即可自动生成目录）"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_txt)
    run._r.append(fld_end)


def page_break(doc):
    doc.add_page_break()


def build():
    doc = Document()
    # 默认样式
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    # 页边距
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)

    # ===== 封面 =====
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("网上国网技术方案 AI 审查系统")
    set_run(run, size=26, bold=True, cn=TITLE_FONT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("软件说明书")
    set_run(run, size=22, bold=True, cn=TITLE_FONT)
    doc.add_paragraph()
    for line in ["版本号：V1.0", "文档日期：2026 年 8 月", "编制单位：网上国网运营中心"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_run(run, size=14, cn=BODY_FONT)
    page_break(doc)

    # ===== 目录 =====
    h(doc, "目录", level=1)
    add_toc(doc)
    page_break(doc)

    # ===== 一、软件概述 =====
    h(doc, "一、软件概述", 1)
    h(doc, "1.1 开发目的", 2)
    para(doc, "网上国网核心系统的技术方案评审，长期依赖评审专家逐份精读方案文本，存在三方面突出问题：一是评审效率低，专家需投入大量时间逐字审阅篇幅较长的技术方案；二是评审口径不一致，不同专家对同类问题的判定标准存在差异；三是隐性风险易漏检，尤其是容量评估、安全防护、接口边界等专业性较强的隐性风险点，仅凭人工经验难以全面覆盖。")
    para(doc, "本软件面向技术方案审核专家，提供一套轻量化的 AI 预审辅助工具。专家提交技术方案（Markdown 格式）后，系统结合内置的审查规则库与审查经验库，自动产出分级审查意见，专家在此基础上复核确认，输出正式评审意见，从而提升评审效率、统一评审口径、降低漏检风险。")
    h(doc, "1.2 软件特点", 2)
    bullet(doc, "轻量化：系统采用单机部署方式，数据全部以本地文件存储，无需数据库与向量库，部署与运维成本低。")
    bullet(doc, "可配置：大模型调用地址、模型名称、密钥等参数通过配置文件管理，可适配外网 API 与内网私有化部署的大模型。")
    bullet(doc, "可量化：内置质量评测体系，采用准确率、召回率等指标持续度量审查质量，支撑系统持续优化。")
    bullet(doc, "可迭代：审查规则库与经验库以文件形式维护，专家可随时人工修改调整，修改后立即生效。")
    h(doc, "1.3 适用范围", 2)
    para(doc, "本软件适用于网上国网核心系统技术方案的评审场景，覆盖能力开放、功能优化、数据共享、消息推送等多种方案类型。第一版以「能力开放 / 三方接入」类技术方案为试点场景，后续可横向扩展至更多场景类型。")
    h(doc, "1.4 软件功能结构", 2)
    para(doc, "本软件由三大功能模块构成，分别是：方案审查模块、数据运营模块、质量评测模块。三个模块共用同一套本地文件数据层，形成「审查 → 台账 → 评测」的完整闭环。")
    table(doc,
          ["模块", "功能说明"],
          [
              ["方案审查", "对技术方案进行 AI 预审，产出分级审查意见，支持结果导出"],
              ["数据运营", "展示微服务台账，以及技术方案与微服务的双向变更关系"],
              ["质量评测", "用准确率 / 召回率指标量化审查质量，辅助持续优化"],
          ],
          widths=[4, 12])
    page_break(doc)

    # ===== 二、运行环境 =====
    h(doc, "二、运行环境", 1)
    h(doc, "2.1 硬件环境", 2)
    table(doc, ["项目", "最低要求", "推荐配置"],
          [
              ["处理器", "双核 CPU", "四核及以上 CPU"],
              ["内存", "4 GB", "8 GB 及以上"],
              ["硬盘", "2 GB 可用空间", "10 GB 及以上"],
          ], widths=[3, 6, 6])
    h(doc, "2.2 软件环境", 2)
    table(doc, ["类别", "名称", "版本要求"],
          [
              ["操作系统", "Windows / Linux", "Windows 10 及以上 / 主流 Linux 发行版"],
              ["运行时", "Python", "3.11 及以上"],
              ["依赖库", "FastAPI、Uvicorn、OpenAI SDK、PyYAML、pdfplumber", "最新稳定版"],
              ["浏览器", "Chrome / Edge 等现代浏览器", "最新稳定版"],
          ], widths=[3, 8, 6])
    h(doc, "2.3 网络环境", 2)
    para(doc, "系统运行需能访问所配置的大模型 API 地址。当采用外网大模型时，需具备外网访问能力；当切换为内网私有化模型时，仅需内网连通即可，业务数据全程不出内网。")
    page_break(doc)

    # ===== 三、安装与部署 =====
    h(doc, "三、安装与部署", 1)
    h(doc, "3.1 环境准备", 2)
    para(doc, "确认操作系统已安装 Python 3.11 及以上版本，并已安装 pip 包管理工具。")
    h(doc, "3.2 依赖安装", 2)
    para(doc, "在项目根目录下执行以下命令安装依赖库：", indent=False)
    p = doc.add_paragraph()
    run = p.add_run("pip install fastapi uvicorn openai pyyaml pdfplumber")
    set_run(run, size=11, cn="Consolas")
    p.paragraph_format.left_indent = Pt(24)
    h(doc, "3.3 系统配置", 2)
    para(doc, "系统配置文件为项目根目录下的 config.yaml，主要配置项如下表：", indent=False)
    table(doc, ["配置项", "说明", "示例值"],
          [
              ["llm.base_url", "大模型调用地址", "https://api.deepseek.com"],
              ["llm.api_key", "大模型访问密钥", "（从环境变量或配置文件注入）"],
              ["llm.model", "模型名称", "deepseek-v4-flash"],
              ["llm.max_tokens", "单次调用最大令牌数", "8192"],
              ["server.host / port", "服务监听地址与端口", "127.0.0.1 / 8000"],
          ], widths=[4, 6, 6])
    h(doc, "3.4 启动系统", 2)
    para(doc, "在项目根目录下执行以下命令启动服务：", indent=False)
    p = doc.add_paragraph()
    run = p.add_run("python -m uvicorn src.app:app --host 127.0.0.1 --port 8000")
    set_run(run, size=11, cn="Consolas")
    p.paragraph_format.left_indent = Pt(24)
    para(doc, "服务启动成功后，打开浏览器访问 http://127.0.0.1:8000 即可进入系统主界面。")
    h(doc, "3.5 停止系统", 2)
    para(doc, "在启动服务的终端窗口按下 Ctrl+C 即可停止服务。")
    page_break(doc)

    # ===== 四、功能模块说明 =====
    h(doc, "四、功能模块说明", 1)
    para(doc, "系统主界面顶部提供三个功能页签：方案审查、数据运营、质量评测，点击页签即可在各功能模块之间切换。")

    h(doc, "4.1 方案审查", 2)
    para(doc, "功能说明：对技术方案进行 AI 预审，产出分级审查意见。")
    para(doc, "操作步骤如下：", indent=False)
    bullet(doc, "第一步：在左侧文本框内粘贴技术方案的 Markdown 内容；也可点击「载入种子方案」按钮载入示例方案。")
    bullet(doc, "第二步：点击「开始审查」按钮。")
    bullet(doc, "第三步：系统自动完成方案解析、规则硬检查、大模型综合审查，并在右侧展示分级审查意见。")
    bullet(doc, "第四步：如需留存，可点击「导出结果」按钮，将审查意见导出为 Markdown 文件。")
    para(doc, "审查意见采用分级呈现方式，第一层为整体结论（通过 / 建议整改 / 不通过）与要点摘要；第二层为逐条问题明细，每条问题包含严重级别（阻塞 / 重要 / 建议）、问题描述、证据定位（引用方案原文关键句）与整改建议。")

    h(doc, "4.2 数据运营", 2)
    para(doc, "功能说明：展示微服务台账，以及「技术方案 ↔ 微服务」的双向变更关系，支撑技术方案的变更统计与运营分析。")
    para(doc, "页面展示内容包括：", indent=False)
    bullet(doc, "统计卡片：微服务总数、技术方案数、变更记录数、被调整微服务数。")
    bullet(doc, "微服务台账表：展示各微服务的名称、所属系统、内外网位置、调整次数、最近调整时间；点击某行可展开该微服务被哪些方案调整过的变更历史。")
    bullet(doc, "技术方案变更列表：展示各方案的标题、场景、状态、结论及变更的微服务；点击某行可展开该方案调整了哪些微服务。")

    h(doc, "4.3 质量评测", 2)
    para(doc, "功能说明：采用准确率、召回率等指标，量化 AI 审查质量，辅助系统持续优化。")
    para(doc, "操作步骤如下：", indent=False)
    bullet(doc, "第一步：切换到「质量评测」页签。")
    bullet(doc, "第二步：点击「运行评测」按钮。")
    bullet(doc, "第三步：系统基于评测集（含专家标注的标准答案）与 AI 审查结果，计算并展示各项指标。")
    para(doc, "页面展示内容包括：核心指标卡片（准确率、召回率、F1、加权召回率）、各维度指标表（按审查维度分别统计 TP / FP / FN 与各项指标）、漏检明细（AI 未检出的问题）、误报明细（AI 误报的问题）。")
    para(doc, "核心指标含义如下表：", indent=False)
    table(doc, ["指标", "含义", "业务解读"],
          [
              ["准确率 Precision", "AI 报出的问题中真问题的比例", "衡量「不误报」"],
              ["召回率 Recall", "专家标注的问题中 AI 检出的比例", "衡量「不漏检」"],
              ["F1", "准确率与召回率的调和平均", "综合衡量审查质量"],
              ["加权召回率", "按严重级别加权的召回率", "突出高风险问题漏检的代价"],
          ], widths=[4, 6, 6])
    page_break(doc)

    # ===== 五、审查规则库与经验库 =====
    h(doc, "五、审查规则库与经验库", 1)
    h(doc, "5.1 审查规则库", 2)
    para(doc, "审查规则库存于 rules 目录，以结构化 JSON 文件维护。第一版规则包为「能力开放 / 三方接入类」规则包，共 24 项检查规则，覆盖 8 个审查维度。")
    table(doc, ["维度", "维度说明", "优先级"],
          [
              ["容量与性能", "限流、扩容、压测、性能影响评估", "P0"],
              ["安全与合规", "加密防护、鉴权闭环、实名信息、数据暴露面", "P0"],
              ["接口与边界", "接口文档、边界逻辑、改动影响、渠道标识", "P0"],
              ["渠道管控", "渠道维度的限流、开关、配额管控", "P1"],
              ["数据同步", "同步机制、异步策略", "P1"],
              ["可配置化", "关键参数可配置", "P1"],
              ["埋点监控", "埋点技术实现方式", "P1"],
              ["文档完整性", "时序图、改造逻辑、变更原因、原型对比", "P1"],
          ], widths=[3, 9, 2])
    h(doc, "5.2 审查经验库", 2)
    para(doc, "审查经验库存于 knowledge 目录，以 Markdown 文件按维度组织，记录高频问题、典型意见摘录与审查要点，供大模型审查时参考。经验库文件可随时人工修改，修改后立即生效。")
    page_break(doc)

    # ===== 六、数据文件与目录结构 =====
    h(doc, "六、数据文件与目录结构", 1)
    para(doc, "系统采用「文件即数据库」的设计，所有数据以本地文件存储，主要目录与文件如下表：", indent=False)
    table(doc, ["目录 / 文件", "用途"],
          [
              ["config.yaml", "大模型与服务配置"],
              ["rules/*.json", "审查规则库"],
              ["knowledge/*.md", "审查经验库"],
              ["data/ledger/services.json", "微服务台账"],
              ["data/ledger/solutions.json", "技术方案与变更记录"],
              ["data/eval/golden_set.json", "评测集（含专家标注标准答案）"],
              ["data/eval/review_records.json", "历史评审意见结构化记录"],
              ["src/", "系统源代码"],
              ["web/index.html", "前端界面"],
              ["scripts/", "辅助脚本（评测、意见单解析等）"],
          ], widths=[8, 8])
    page_break(doc)

    # ===== 七、常见问题 =====
    h(doc, "七、常见问题", 1)
    h(doc, "7.1 服务启动失败", 2)
    para(doc, "请检查：Python 版本是否满足要求；依赖库是否安装完整；config.yaml 配置是否正确；端口是否被占用。")
    h(doc, "7.2 审查结果为空", 2)
    para(doc, "请检查所配置的大模型 API 地址与密钥是否有效，网络是否可达。")
    h(doc, "7.3 评测页面无数据", 2)
    para(doc, "请点击「运行评测」按钮触发评测；若仍无数据，请检查 data/eval 目录下评测集文件是否存在。")
    page_break(doc)

    # ===== 八、技术架构说明 =====
    h(doc, "八、技术架构说明", 1)
    para(doc, "系统采用前后端分离的单机架构，后端基于 Python FastAPI 框架，前端为原生 HTML / JavaScript 单页应用。整体分为四层：")
    bullet(doc, "接入层：Web 界面，提供方案审查、数据运营、质量评测三个功能页签。")
    bullet(doc, "服务层：方案解析器、规则引擎、经验检索、大模型生成、审查编排、台账数据层、评测器。")
    bullet(doc, "文件数据层：审查规则库、审查经验库、台账数据、评测数据，均以本地 JSON / Markdown 文件存储。")
    bullet(doc, "模型层：统一的大模型客户端封装，采用 OpenAI 兼容协议，可灵活切换模型厂商。")
    page_break(doc)

    # ===== 附录 A：审查规则清单 =====
    import json
    h(doc, "附录 A：审查规则清单", 1)
    para(doc, "第一版规则包共包含 24 项检查规则，完整清单如下表所示：", indent=False)
    dim_names = {"capacity": "容量与性能", "security": "安全与合规", "interface": "接口与边界",
                 "channel": "渠道管控", "data_sync": "数据同步", "config": "可配置化",
                 "monitoring": "埋点监控", "documentation": "文档完整性"}
    sev_names = {"blocker": "阻塞", "major": "重要", "minor": "建议"}
    with open(r"C:\Users\zhito\WorkBuddy\2026-08-17-09-25-53\rules\capability-open-access-rules.json", "r", encoding="utf-8") as f:
        rules = json.load(f)["rule_pack"]["checks"]
    rows = [[c["id"], dim_names.get(c["dimension"], c["dimension"]), c["name"], sev_names.get(c.get("severity"), c.get("severity"))] for c in rules]
    table(doc, ["规则 ID", "维度", "检查项", "严重级别"], rows, widths=[2.5, 3.5, 7.5, 2.5])
    page_break(doc)

    # ===== 附录 B：术语表 =====
    h(doc, "附录 B：术语表", 1)
    table(doc, ["术语", "说明"],
          [
              ["AI 预审", "由 AI 对技术方案进行初步审查，产出结构化意见供专家复核"],
              ["分级审查意见", "包含整体结论、逐条风险、证据定位、整改建议的结构化意见"],
              ["准确率（Precision）", "AI 报出的问题中真问题的比例"],
              ["召回率（Recall）", "专家标注的问题中 AI 检出的比例"],
              ["规则硬检查", "由程序按确定性规则执行的检查，不消耗大模型资源"],
              ["证据定位", "将审查意见锚定到方案原文具体章节或语句"],
          ], widths=[5, 11])

    doc.save(OUT)
    print("已生成:", OUT)


if __name__ == "__main__":
    build()
